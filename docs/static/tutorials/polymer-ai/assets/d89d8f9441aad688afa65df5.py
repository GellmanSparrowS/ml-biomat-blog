"""L16: loading history, learned states, and experimental cyclic curves.
Run: python example.py
Install: numpy scipy scikit-learn torch matplotlib openpyxl requests
The synthetic teacher parameters are specified below. Experimental source data
are read separately; they are not training labels for this small-strain model.
"""
from pathlib import Path
import json,copy,time
import numpy as np
from scipy.optimize import nnls
from sklearn.linear_model import Ridge
import torch
from torch import nn
torch.set_num_threads(2)
HERE=Path.cwd()
OUT=HERE/"history_outputs";OUT.mkdir(exist_ok=True)
DT=.05;T=np.arange(401)*DT;E0=6.;E=np.array([12.,8.]);TAU=np.array([.3,3.])
rng=np.random.default_rng(1616)

def pulse(onset,rise,hold,fall,amplitude):
    knots=np.array([0,onset,onset+rise,onset+rise+hold,onset+rise+hold+fall,20.])
    return np.interp(T,knots,[0,0,amplitude,amplitude,0,0])

def duration(lo,hi):return round(rng.uniform(lo,hi)/DT)*DT

strain=[];records=[]
for k in range(304):
    amp=rng.uniform(.02,.1);on=duration(.2,.9)
    kind='single' if k<240 else 'double' if k<272 else 'fast'
    rise=duration(.4,3.2) if kind!='fast' else duration(.1,.25)
    fall=duration(.4,3.2) if kind!='fast' else duration(.1,.25)
    hold=duration(.4,3.)
    eps=pulse(on,rise,hold,fall,amp)
    if kind=='double':
        rise=duration(.4,1.2);fall=duration(.4,1.2);hold=duration(.4,1.2)
        eps=pulse(on,rise,hold,fall,amp)+pulse(10+on,rise,hold,fall,.7*amp)
    strain.append(eps);records.append(dict(id=k,kind=kind,amplitude=amp,onset=on,rise=rise,hold=hold,fall=fall))
strain=np.array(strain);rate=np.diff(strain,axis=1)/DT

def integrate(eps,e0,moduli,times):
    n,steps=eps.shape;q=np.zeros((n,len(times)));out=np.zeros_like(eps)
    states=np.zeros((n,steps,len(times)));work=np.zeros_like(eps);diss=np.zeros_like(eps)
    a=np.exp(-DT/times);one=-np.expm1(-DT/times)
    for j in range(1,steps):
        r=(eps[:,j]-eps[:,j-1])/DT;c=moduli*times*r[:,None];b=q-c
        iq=c*DT+b*times*one
        iq2=c*c*DT+2*c*b*times*one+b*b*times*.5*(-np.expm1(-2*DT/times))
        work[:,j]=work[:,j-1]+e0*.5*(eps[:,j]**2-eps[:,j-1]**2)+(r[:,None]*iq).sum(1)
        diss[:,j]=diss[:,j-1]+(iq2/(moduli*times)).sum(1)
        q=a*q+moduli*times*one*r[:,None];states[:,j]=q
        out[:,j]=e0*eps[:,j]+q.sum(1)
    energy=e0*.5*eps**2+(states**2/(2*moduli)).sum(2)
    return out,states,work,diss,energy

truth,states,work,diss,energy=integrate(strain,E0,E,TAU)
assert np.max(abs(work-energy-diss))<1e-12
assert np.min(np.diff(diss,axis=1))>-1e-13
ids=rng.permutation(240);train=ids[:144];val=ids[144:192];test=ids[192:]
splits=dict(train=train.tolist(),validation=val.tolist(),test=test.tolist(),double=list(range(240,272)),fast=list(range(272,304)))
noise=rng.normal(0,.008,truth.shape);noise[:,0]=0;observed=truth+noise
features=np.stack([strain/.1,np.column_stack([np.zeros(304),np.diff(strain,axis=1)])/.1],axis=2)
# Current-strain model receives no time or history state.
poly=np.stack([strain/.1,(strain/.1)**2,(strain/.1)**3],axis=2)
ridge=Ridge(alpha=1e-6).fit(poly[train].reshape(-1,3),observed[train].ravel())
predictions={'strain_only':ridge.predict(poly.reshape(-1,3)).reshape(truth.shape)}

# A fixed positive relaxation bank learns the nonnegative strength of each mode.
bank=np.geomspace(.05,12.,12);history=np.zeros((304,401,13));history[:,:,0]=strain
for m,tau in enumerate(bank):
    h=np.zeros(304);a=np.exp(-DT/tau);b=tau*(-np.expm1(-DT/tau))
    for j in range(1,401):
        h=a*h+b*rate[:,j-1];history[:,j,m+1]=h
weights,residual=nnls(history[train].reshape(-1,13),observed[train].ravel())
predictions['positive_state']=history@weights

class FeedForward(nn.Module):
    def __init__(self):
        super().__init__();self.layers=nn.Sequential(nn.Linear(2,16),nn.Tanh(),nn.Linear(16,16),nn.Tanh(),nn.Linear(16,1))
    def forward(self,x):return self.layers(x).squeeze(-1)

class Recurrent(nn.Module):
    def __init__(self):
        super().__init__();self.gru=nn.GRU(2,16,batch_first=True);self.readout=nn.Linear(16,1)
    def forward(self,x):
        h,_=self.gru(x);return self.readout(h).squeeze(-1)

xt=torch.tensor(features,dtype=torch.float32);yt=torch.tensor(observed,dtype=torch.float32)
models={};histories={};seeds=[16,26,36]
start=time.monotonic()
for name,cls in [('instantaneous',FeedForward),('gru',Recurrent)]:
    seed_preds=[];models[name]=[];histories[name]=[]
    for seed in seeds:
        torch.manual_seed(seed);model=cls();optimizer=torch.optim.Adam(model.parameters(),lr=.01)
        best=float('inf');best_state=None;curve=[];best_epoch=None
        for epoch in range(1,121):
            model.train();optimizer.zero_grad();prediction=model(xt[train]);loss=((prediction-yt[train])**2).mean()
            loss.backward();nn.utils.clip_grad_norm_(model.parameters(),1.);optimizer.step()
            if epoch%5==0:
                model.eval()
                with torch.no_grad():v=float(((model(xt[val])-yt[val])**2).mean())
                curve.append([epoch,float(loss.detach()),v])
                if v<best:best=v;best_state=copy.deepcopy(model.state_dict());best_epoch=epoch
        model.load_state_dict(best_state);model.eval()
        with torch.no_grad():p=model(xt).numpy().astype(float)
        seed_preds.append(p)
        models[name].append(dict(seed=seed,best_epoch=best_epoch,validation_mse=best,
            parameters=sum(v.numel() for v in model.parameters()),weights={k:v.tolist() for k,v in best_state.items()}))
        histories[name].append(curve)
        print(f'{name} seed{seed}: epoch{best_epoch}, validation RMSE{best**.5:.6f}',flush=True)
    predictions[name]=np.mean(seed_preds,axis=0)

def rmse(a,b):return float(np.sqrt(np.mean((a-b)**2)))
scores={name:{split:dict(clean_RMSE_kPa=rmse(p[ix],truth[ix]),observed_RMSE_kPa=rmse(p[ix],observed[ix]))
              for split,ix in splits.items() if split!='train'} for name,p in predictions.items()}
result=dict(dt=DT,teacher=dict(E0_kPa=E0,E_kPa=E.tolist(),tau_s=TAU.tolist()),records=records,splits=splits,
    scores=scores,models=models,history=histories,positive_state=dict(tau_s=bank.tolist(),moduli_kPa=weights.tolist()),
    strain_only=dict(coef=ridge.coef_.tolist(),intercept=float(ridge.intercept_)),
    checks=dict(energy_balance_max=float(np.max(abs(work-energy-diss))),minimum_dissipation_increment=float(np.min(np.diff(diss,axis=1)))))

# Whole-path scores and a second window that includes the active loading process.
active={}
for name,p in predictions.items():
    active[name]={}
    for split in ['test','double','fast']:
        values=[]
        for k in splits[split]:
            rec=records[k]
            end=rec['onset']+rec['rise']+rec['hold']+rec['fall']+3+(10 if rec['kind']=='double' else 0)
            mask=T<=min(20,end)
            values.append(rmse(p[k,mask],truth[k,mask]))
        active[name][split]=float(np.mean(values))
result['active_window_mean_path_RMSE_kPa']=active
result['data_definition']=dict(seed=1616,noise_sd_kPa=.008,time_step_s=DT,
    train_unit='complete loading protocol',units='stress kPa; energy density kJ/m3',
    teacher='chosen small-strain Maxwell solid, constant temperature, no damage',
    architecture='PyTorch GRU: reset gate applied after hidden affine transform')
zero_input=torch.zeros((1,2001,2));zero={}
for name,cls in [('instantaneous',FeedForward),('gru',Recurrent)]:
    rest=[]
    for item in models[name]:
        model=cls();model.load_state_dict({k:torch.tensor(v) for k,v in item['weights'].items()});model.eval()
        with torch.no_grad():rest.append(model(zero_input).numpy()[0])
    y=np.mean(rest,axis=0)
    zero[name]=dict(initial_kPa=float(y[0]),after100s_kPa=float(y[-1]))
result['zero_input_100s']=zero

# Read the unmodified publisher workbook. A local copy avoids repeated downloads.
import hashlib,csv
import requests
from openpyxl import load_workbook
SOURCE='https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-023-38280-4/MediaObjects/41467_2023_38280_MOESM7_ESM.xlsx'
SOURCE_SHA='788fa83caa5b2f835a63bab45f06107c17fe1c4f2e78acefb195e48334bbd809'
source=HERE/'peptide_source_data.xlsx'
if not source.exists():
    temporary=source.with_suffix('.xlsx.part')
    for attempt in range(4):
        offset=temporary.stat().st_size if temporary.exists() else 0
        try:
            with requests.get(SOURCE,headers={'Range':f'bytes={offset}-'} if offset else {},stream=True,timeout=(20,90)) as response:
                response.raise_for_status()
                mode='ab' if offset and response.status_code==206 else 'wb'
                with temporary.open(mode) as stream:
                    for chunk in response.iter_content(1024*1024):stream.write(chunk)
            assert hashlib.sha256(temporary.read_bytes()).hexdigest()==SOURCE_SHA
            temporary.replace(source);break
        except (requests.RequestException,AssertionError):
            if attempt==3:raise
assert hashlib.sha256(source.read_bytes()).hexdigest()==SOURCE_SHA
workbook=load_workbook(source,read_only=True,data_only=True)
curves=[]
for sheet,columns in [('Figure 4c',[(1,'PAM',1),(4,'p-Pep',1),(7,'p-Pep/Cu2+',1)]),
                      ('Figure 4e',[(25,'p-Pep/Cu2+',1),(45,'p-Pep/Cu2+',100)])]:
    pairs=[[] for _ in columns]
    for row in workbook[sheet].iter_rows(min_row=6,values_only=True):
        for j,(col,material,cycle) in enumerate(columns):
            x,y=row[col:col+2]
            if isinstance(x,(int,float)) and isinstance(y,(int,float)):pairs[j].append([float(x),float(y)])
    for points,(col,material,cycle) in zip(pairs,columns):
        a=np.asarray(points)
        curves.append(dict(sheet=sheet,material=material,cycle=cycle,columns=[col+1,col+2],
            points=a,n=len(a),peak_MPa=float(a[:,1].max()),
            ordered_curve_integral_MPa=float(np.trapezoid(a[:,1],a[:,0]))))
summary=[]
for row in workbook['Supp. Fig. 21a'].iter_rows(min_row=6,values_only=True):
    if isinstance(row[4],(int,float)):
        summary.append([int(row[4]),float(row[5]),float(row[11])])
workbook.close();summary=np.array(summary)
assert len(summary)==100 and summary[-1,0]==100
result['experiment']=dict(doi='10.1038/s41467-023-38280-4',source_url=SOURCE,sha256=SOURCE_SHA,
    curves=[{k:v for k,v in c.items() if k!='points'} for c in curves],
    summary_sheet='Supp. Fig. 21a',summary_columns=[5,6,12],
    copper_cycle100_peak_percent=float(summary[-1,1]),
    copper_cycle100_dissipation_percent=float(summary[-1,2]),
    copper_integral_retention_percent=100*curves[-1]['ordered_curve_integral_MPa']/curves[-2]['ordered_curve_integral_MPa'],
    material='peptide fibres in a polyacrylamide matrix; with/without copper coordination',
    sequence_note='Repeated cycles from the same specimen are not independent specimen replicates')
with (OUT/'experimental_curves.csv').open('w',encoding='utf-8',newline='') as stream:
    writer=csv.writer(stream);writer.writerow(['sheet','material','cycle','point_order','strain_mm_per_mm','stress_MPa'])
    for c in curves:
        writer.writerows([c['sheet'],c['material'],c['cycle'],i,*xy] for i,xy in enumerate(c['points']))
(OUT/'experimental_cycle_summary.json').write_text(json.dumps(summary.tolist(),indent=2)+'\n',encoding='utf-8')
(OUT/'results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
samples=dict(records=records,splits=splits,time_s=T.tolist(),strain=strain.tolist(),
    clean_stress_kPa=truth.tolist(),observed_stress_kPa=observed.tolist())
(OUT/'samples.json').write_text(json.dumps(samples,separators=(',',':'))+'\n',encoding='utf-8')
np.savez_compressed(OUT/'predictions.npz',**predictions,states=states,work=work,dissipation=diss,energy=energy)

# Consistent axes, units, and panel order for the offline HTML and notebook.
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,
    'axes.spines.right':False,'axes.labelcolor':'#30453f','text.color':'#243d35',
    'axes.edgecolor':'#bac7c0','xtick.color':'#53675e','ytick.color':'#53675e',
    'svg.hashsalt':'L16-history','figure.facecolor':'#fffef9','axes.facecolor':'#fffef9'})
names=['strain_only','instantaneous','gru','positive_state']
labels=['Strain only','Instantaneous MLP','GRU history','Positive modes']
colors=['#a2937b','#b57552','#387da0','#287b63']
def style(ax,title,xlabel='Time (s)',ylabel='Stress (kPa)'):
    ax.set_title(title,loc='left',fontsize=12,pad=13,weight='bold');ax.set_xlabel(xlabel);ax.set_ylabel(ylabel)
    ax.grid(axis='y',alpha=.18);ax.set_axisbelow(True)
demo=np.array([pulse(.5,d,1,d,.08) for d in [.5,3]])
demo_stress,_,demo_work,demo_diss,demo_energy=integrate(demo,E0,E,TAU)
def panel(i,fig,spec):
    if i==0:
        grid=spec.subgridspec(2,1,hspace=.08,height_ratios=[1,2]);axes=[fig.add_subplot(grid[j]) for j in range(2)]
        for j,(label,color) in enumerate([('0.5 s ramp','#bc8452'),('3 s ramp','#387e94')]):
            axes[0].plot(T,demo[j]*100,color=color,label=label)
            axes[1].plot(T,demo_stress[j],color=color,label=label,lw=2)
        style(axes[0],'A  Same amplitude, different speed','','Strain (%)');axes[0].tick_params(labelbottom=False)
        style(axes[1],'',ylabel='Stress (kPa)');axes[0].legend(frameon=False,ncol=2,fontsize=9)
        for a in axes:a.set_xlim(0,12)
        axes[1].axhline(0,color='#b8c7bb',lw=.8)
        return
    ax=fig.add_subplot(spec)
    if i==1:
        k=int(test[0]);ax.plot(strain[k]*100,truth[k],color='#223b34',lw=2.7,label='Reference')
        for name,label,color in zip(names,labels,colors):ax.plot(strain[k]*100,predictions[name][k],color=color,lw=1.3,ls='--',label=label)
        style(ax,'B  A complete held-out loop','Strain (%)');ax.legend(frameon=False,fontsize=8,loc='upper left')
        ax.text(.98,.03,f'Protocol {k}',transform=ax.transAxes,ha='right',fontsize=9)
    elif i==2:
        k=240;ax.plot(T,truth[k],color='#223b34',lw=2.7,label='Reference')
        for name,label,color in zip(names,labels,colors):ax.plot(T,predictions[name][k],color=color,lw=1.3,ls='--',label=label)
        style(ax,'C  Two pulses: unseen history');ax.set_xlim(0,20)
        ax.legend(frameon=False,fontsize=8,ncol=2,loc='upper right');ax.axhline(0,color='#b8c7bb',lw=.8)
    elif i==3:
        x=np.arange(3)
        for j,(name,label,color) in enumerate(zip(names,labels,colors)):
            values=[scores[name][s]['clean_RMSE_kPa'] for s in ['test','double','fast']]
            positions=x+(j-1.5)*.18
            ax.bar(positions,values,.16,color=color,label=label)
            ax.scatter(positions,[active[name][s] for s in ['test','double','fast']],s=17,facecolor='white',edgecolor='#263e36',zorder=5)
        ax.set_xticks(x,['Single / 48','Double / 32','Fast / 32']);ax.set_yscale('log');ax.set_ylim(.0003,.65)
        style(ax,'D  Hold out complete protocols','Protocol family / count','RMSE (kPa, log scale)')
        ax.legend(frameon=False,fontsize=8,ncol=2,loc='upper right')
        ax.text(.02,.03,'Bars: full 20 s  ·  circles: active window',transform=ax.transAxes,fontsize=9,bbox=dict(facecolor='#fffef9',edgecolor='none',alpha=.9,pad=3))
    elif i==4:
        for y,label,color,ls in [(demo_work[0],'Cumulative work','#384f73','-'),
                (demo_energy[0],'Stored energy','#b78351','--'),(demo_diss[0],'Dissipated energy','#287b63','-')]:
            ax.plot(T,y,color=color,lw=2,ls=ls,label=label)
        style(ax,'E  Where did the work go?',ylabel='Energy density (kJ/m³)');ax.set_xlim(0,12)
        ax.legend(frameon=False,fontsize=9);ax.text(.98,.12,'Work = stored + dissipated',transform=ax.transAxes,ha='right',fontsize=9)
    else:
        nonzero=weights[1:]>1e-8
        ax.vlines(bank[nonzero],0,weights[1:][nonzero],color='#287b63',lw=4,label='Fitted positive modes')
        ax.scatter(bank[nonzero],weights[1:][nonzero],s=40,color='#287b63',zorder=3)
        ax.vlines(TAU,0,E,color='#b78351',lw=2,linestyle='--',label='Reference modes')
        ax.scatter(TAU,E,marker='D',s=44,facecolor='white',edgecolor='#b78351',zorder=4)
        ax.set_xscale('log');ax.set_xlim(.045,14);ax.set_ylim(0,15)
        style(ax,'F  A learned spectrum of time scales','Relaxation time (s, log scale)','Mode modulus (kPa)')
        ax.legend(frameon=False,fontsize=9,loc='upper right')
        ax.text(.03,.04,f'Equilibrium modulus: {weights[0]:.3f} kPa\nReference: 6.000 kPa',transform=ax.transAxes,fontsize=9)
keys=['speed','loop','double','scores','energy','spectrum']
for i,key in enumerate(keys):
    fig=plt.figure(figsize=(8.2,5.1),layout='constrained');panel(i,fig,fig.add_gridspec(1,1)[0])
    fig.savefig(OUT/f'{key}.svg',metadata={'Date':None});plt.close(fig)
fig=plt.figure(figsize=(16.4,15.3),layout='constrained');grid=fig.add_gridspec(3,2,hspace=.09,wspace=.05)
for i in range(6):panel(i,fig,grid[i//2,i%2])
fig.savefig(OUT/'history_learning.svg',metadata={'Date':None})
fig.savefig(OUT/'history_learning.png',dpi=160,metadata={'Software':'L16 reproducible example'});plt.close(fig)
for j,key in enumerate(['materials','cycles','retention']):
    fig,ax=plt.subplots(figsize=(8.2,5.1),layout='constrained')
    if j==0:
        for c,color in zip(curves[:3],['#969a83','#b78351','#287b63']):
            a=c['points'];ax.plot(a[:,0],a[:,1],color=color,lw=1.8,label=c['material'])
        style(ax,'I  Three material formulations','Strain (mm/mm)','Stress (MPa)');ax.legend(frameon=False)
        ax.text(.03,.92,'Source: Figure 4c',transform=ax.transAxes,fontsize=9)
    elif j==1:
        for c,color in zip(curves[3:],['#287b63','#b78351']):
            a=c['points'];ax.plot(a[:,0],a[:,1],color=color,lw=1.7,label=f"Cycle {c['cycle']}")
        style(ax,'II  Peptide / PAM / copper composite','Strain (mm/mm)','Stress (MPa)');ax.legend(frameon=False)
        ax.text(.03,.92,'Source: Figure 4e',transform=ax.transAxes,fontsize=9)
    else:
        ax.plot(summary[:,0],summary[:,1],color='#287b63',lw=2,label='Peak stress')
        ax.plot(summary[:,0],summary[:,2],color='#b78351',lw=2,label='Dissipation')
        style(ax,'III  Repeated use of the same specimen','Cycle number','Relative to first cycle (%)')
        ax.set_ylim(85,102);ax.legend(frameon=False,loc='lower left')
        ax.text(.98,.94,f'Cycle 100: {summary[-1,1]:.2f}% / {summary[-1,2]:.2f}%',transform=ax.transAxes,ha='right',fontsize=9)
    fig.savefig(OUT/f'experimental_{key}.svg',metadata={'Date':None});plt.close(fig)
print('Teacher: chosen small-strain model; experimental curves are a separate companion')
print('144 train / 48 validation / 48 test; 32 double and 32 fast extra histories')
print('Model                    single       double       fast   (clean RMSE, kPa)')
for name,label in zip(names,labels):
    values=[scores[name][s]['clean_RMSE_kPa'] for s in ['test','double','fast']]
    print(f'{label:22s} '+''.join(f'{v:12.6f}' for v in values))
print(f"Maximum exact energy-balance residual: {result['checks']['energy_balance_max']:.2e} kJ/m3")
print(f"Experimental Cu composite integral retention: {result['experiment']['copper_integral_retention_percent']:.4f}%")
print('Six model panels, three experimental panels and complete numerical results: history_outputs/')
