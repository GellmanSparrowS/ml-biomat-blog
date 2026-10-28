"""Silk processing: reproduce a four-condition, group-held-out experiment.

Install: pip install numpy matplotlib openpyxl requests scikit-learn
Source: Frydrych et al. (2019), doi:10.1038/s41598-019-51589-9.
Original workbook: doi:10.5281/zenodo.3457387 (read without modifying it).
"""
from pathlib import Path
import hashlib
import json
import numpy as np
import requests
from openpyxl import load_workbook
from sklearn.preprocessing import StandardScaler
from sklearn.linear_model import Ridge
from sklearn.ensemble import RandomForestRegressor
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

OUT = Path("silk_outputs")
OUT.mkdir(exist_ok=True)
DATA = Path("Data.xlsx")
URL = "https://zenodo.org/records/3457387/files/Data.xlsx?download=1"
SHA = "ae467d3711b0e4c78a1c66c368ebb1d94a9c64070f1f60c87f7a57e0e44456f6"
if not DATA.exists():
    response = requests.get(URL, timeout=90)
    response.raise_for_status()
    assert hashlib.sha256(response.content).hexdigest() == SHA
    DATA.write_bytes(response.content)
assert hashlib.sha256(DATA.read_bytes()).hexdigest() == SHA

# One row of the learning table is one complete processing condition.
# Point clouds contain 27 records per condition in the supplied workbook.
wb = load_workbook(DATA, read_only=True, data_only=True)
sheet = wb["Figure 5A - Tensile"]
rows = list(sheet.iter_rows(values_only=True))
draw = np.array([7.5, 10., 12.5, 15.])
beta = np.array([.29, .33, .35, .36])  # independent FTIR, paper Table 1
area = np.array([586., 311., 186., 102.])  # um², descriptive only
observations, curves, groups = [], [], []
for k in range(4):
    points, curve = [], []
    for row_number, row in enumerate(rows[2:], 3):
        strain, stress = row[4*k+2:4*k+4]
        if isinstance(strain, (int, float)) and isinstance(stress, (int, float)):
            observations.append(dict(condition=k, source_row=row_number,
                source_columns=[4*k+3, 4*k+4], draw_ratio=float(draw[k]),
                fracture_stress_MPa=float(stress), strain_percent=float(strain)))
            points.append([stress, strain])
        strain, stress = row[4*k:4*k+2]
        if isinstance(strain, (int, float)) and isinstance(stress, (int, float)):
            curve.append([strain, stress])
    points = np.asarray(points)
    assert len(points) == 27 and np.isfinite(points).all()
    groups.append(dict(draw_ratio=float(draw[k]), n=len(points),
        beta_fraction=float(beta[k]), mean_area_um2=float(area[k]),
        mean=points.mean(0).tolist(), sample_sd=points.std(0, ddof=1).tolist()))
    curves.append(np.asarray(curve, dtype=float))
wb.close()
X = np.column_stack([draw, beta])
Y = np.array([g["mean"] for g in groups])  # [MPa, percent]
SD = np.array([g["sample_sd"] for g in groups])
names = ["feedstock_mean", "draw_ridge", "draw_beta_ridge", "draw_forest"]
labels = ["Feedstock mean", "Draw: ridge", "Draw + structure", "Draw: forest"]
colors = ["#929c99", "#347a8c", "#247765", "#b87c4f"]
predictions = {name: np.zeros_like(Y) for name in names}
folds, fitted = [], []
ridge_check = 0.
for test in range(4):
    train = np.array([i for i in range(4) if i != test])
    assert test not in train
    predictions["feedstock_mean"][test] = Y[train].mean(0)
    record = dict(test_condition=test, train_conditions=train.tolist(), models={})
    models = {}
    for name, columns in [("draw_ridge", [0]), ("draw_beta_ridge", [0, 1])]:
        scaler = StandardScaler().fit(X[train][:, columns])
        xt = scaler.transform(X[train][:, columns])
        xe = scaler.transform(X[[test]][:, columns])
        model = Ridge(alpha=1.).fit(xt, Y[train])
        predictions[name][test] = model.predict(xe)[0]
        center = Y[train].mean(0)
        coefficient = np.linalg.solve(xt.T @ xt + np.eye(len(columns)),
                                      xt.T @ (Y[train] - center))
        independent = (xe @ coefficient + center)[0]
        ridge_check = max(ridge_check, float(np.max(abs(independent - predictions[name][test]))))
        record["models"][name] = dict(columns=columns,
            mean=scaler.mean_.tolist(), scale=scaler.scale_.tolist(),
            coef=model.coef_.tolist(), intercept=model.intercept_.tolist())
        models[name] = (scaler, model)
    forests = []
    for seed in [15, 25, 35]:
        # Joint two-output squared-error trees; outputs retain MPa and percent.
        # One input feature means bootstrap sampling supplies the randomness.
        forest = RandomForestRegressor(n_estimators=100, max_depth=2,
            min_samples_leaf=1, max_features=1., bootstrap=True,
            random_state=seed, n_jobs=1).fit(X[train, :1], Y[train])
        forests.append(forest)
    seed_predictions = np.array([m.predict(X[[test], :1])[0] for m in forests])
    predictions["draw_forest"][test] = seed_predictions.mean(0)
    record["models"]["draw_forest"] = dict(seeds=[15, 25, 35],
        predictions=seed_predictions.tolist(), n_estimators=100, max_depth=2)
    models["draw_forest"] = forests
    folds.append(record)
    fitted.append(models)
assert ridge_check < 1e-10

def score(predicted, observed):
    error = predicted - observed
    return dict(RMSE=np.sqrt(np.mean(error**2, axis=0)).tolist(),
                MAE=np.mean(abs(error), axis=0).tolist())

raw_y = np.array([[r["fracture_stress_MPa"], r["strain_percent"]] for r in observations])
raw_group = np.array([r["condition"] for r in observations])
scores = {name: dict(condition_mean=score(p, Y),
    individual=score(p[raw_group], raw_y),
    interpolation=score(p[[1, 2]], Y[[1, 2]]),
    extrapolation=score(p[[0, 3]], Y[[0, 3]])) for name, p in predictions.items()}
result = dict(source=dict(paper="10.1038/s41598-019-51589-9",
    data="10.5281/zenodo.3457387", sha256=SHA, sheet=sheet.title),
    groups=groups, targets=["fracture_stress_MPa", "strain_percent"],
    predictions={k: v.tolist() for k, v in predictions.items()}, scores=scores,
    folds=folds, checks=dict(independent_ridge_max_error=ridge_check,
        raw_records=len(observations), conditions=4,
        representative_curve_points=[len(c) for c in curves]),
    protocol=dict(split="leave one complete draw condition out",
        alpha=1., structure_input="independent post-characterization FTIR fraction",
        forest_loss="joint squared error in MPa and percent; target units retained",
        measurement="50% RH, ambient temperature, 10 mm gauge, 20%/min strain rate"))
for filename, value in [("results.json", result), ("samples.json", observations),
                         ("curves.json", [c.tolist() for c in curves])]:
    (OUT / filename).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

# All panels use the same executed data. SVG is deterministic and editable.
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.labelcolor": "#30453f", "text.color": "#243d35",
    "axes.edgecolor": "#bac7c0", "xtick.color": "#53675e",
    "ytick.color": "#53675e", "svg.hashsalt": "L15-silk",
    "figure.facecolor": "#fffef9", "axes.facecolor": "#fffef9"})
condition_colors = ["#bc8b58", "#699a89", "#3b8393", "#334f73"]

def style(ax, title, xlabel="Draw ratio", ylabel=None):
    ax.set_title(title, loc="left", fontsize=12, pad=14, weight="bold")
    ax.set_xlabel(xlabel)
    if ylabel:
        ax.set_ylabel(ylabel)
    ax.grid(axis="y", alpha=.18)
    ax.set_axisbelow(True)

def panel(i, fig, spec):
    if i in [1, 4, 5]:
        inner = spec.subgridspec(1, 2, wspace=.08)
        axes = [fig.add_subplot(inner[j]) for j in range(2)]
    else:
        axes = [fig.add_subplot(spec)]
    ax = axes[0]
    if i == 0:
        for k, curve in enumerate(curves):
            ax.plot(curve[:, 0], curve[:, 1], color=condition_colors[k], lw=1.6,
                    label=f"{draw[k]:g}x")
            v = raw_y[raw_group == k]
            ax.scatter(v[:, 1], v[:, 0], color=condition_colors[k], s=18, alpha=.45)
        style(ax, "A  Curves and fracture records", "Engineering strain (%)", "Engineering stress (MPa)")
        ax.set(xlim=(0, 85), ylim=(0, 510))
        ax.legend(ncol=4, frameon=False, fontsize=9, loc="upper right")
        ax.text(.98, .04, "4 representative curves\n108 fracture records", transform=ax.transAxes,
                ha="right", va="bottom", fontsize=9)
    elif i == 1:
        for a, value, sd, title, unit, color in zip(axes, [area, beta*100],
                [[118, 56, 31, 20], [2, 2, 2, 1]],
                ["B  Cross-section", "FTIR structure"],
                ["Mean area (µm²)", "Beta-sheet fraction (%)"], ["#347a8c", "#247765"]):
            a.errorbar(draw, value, yerr=sd, fmt="o-", color=color, capsize=3)
            style(a, title, ylabel=unit)
            a.set_xticks(draw, ["7.5", "10", "12.5", "15"])
            a.tick_params(axis="x", labelsize=8)
        axes[0].text(.96, .92, "n = 120", transform=axes[0].transAxes, ha="right", fontsize=9)
        axes[1].text(.06, .05, "n = 3\nMean ± SD", transform=axes[1].transAxes, fontsize=9)
    elif i in [2, 3]:
        target = i-2
        ax.errorbar(draw, Y[:, target], yerr=SD[:, target], fmt="o", color="#263d35",
                    capsize=4, ms=6, label="Observed mean ± SD", zorder=5)
        for name, label, color, marker in zip(names, labels, colors, ["v", "s", "D", "^"]):
            ax.plot(draw, predictions[name][:, target], marker=marker, color=color,
                    lw=1.3, ms=4, label=label, linestyle="--")
        style(ax, ["C  Predict fracture stress", "D  Predict strain at break"][target],
              ylabel=["Engineering stress (MPa)", "Engineering strain (%)"][target])
        ax.set_xlim(6.9, 15.6)
        ax.set_ylim((130, 470) if target == 0 else (0, 79))
        ax.set_xticks(draw)
        ax.legend(frameon=False, fontsize=8, ncol=2, loc="upper left")
    elif i == 4:
        for target, a in enumerate(axes):
            values = np.array([scores[n]["condition_mean"]["RMSE"][target] for n in names])
            a.barh(np.arange(4), values, color=colors, height=.58)
            a.set_yticks(range(4), ["Mean", "Draw", "Draw + β", "Forest"] if target == 0 else [])
            a.invert_yaxis()
            a.set_xlim(0, values.max()*1.3)
            for j, v in enumerate(values):
                a.text(v+values.max()*.025, j, f"{v:.1f}", va="center", fontsize=9)
            style(a, ["E  Stress error", "Strain error"][target],
                  ["RMSE (MPa)", "RMSE (percentage points)"][target])
    else:
        grid = np.linspace(6, 17, 301)[:, None]
        for test, a in zip([1, 3], axes):
            train = np.array([j for j in range(4) if j != test])
            scaler, model = fitted[test]["draw_ridge"]
            ridge = model.predict(scaler.transform(grid))[:, 0]
            forest = np.mean([m.predict(grid)[:, 0] for m in fitted[test]["draw_forest"]], axis=0)
            a.axvspan(draw[train].min(), draw[train].max(), color="#e9efdf", zorder=0)
            a.plot(grid[:, 0], ridge, color=colors[1], label="Ridge")
            a.plot(grid[:, 0], forest, color=colors[3], label="Forest")
            a.scatter(draw[train], Y[train, 0], color="#263d35", s=25, zorder=5)
            a.scatter(draw[test], Y[test, 0], marker="D", facecolor="white", edgecolor="#263d35", s=45, zorder=6)
            style(a, "F  Hold out 10x" if test == 1 else "Hold out 15x",
                  ylabel="Stress (MPa)" if test == 1 else None)
            a.set(xlim=(6, 17), ylim=(150, 410), xticks=[7.5, 12.5, 15])
            a.tick_params(axis="x", labelsize=8)
            a.legend(frameon=False, fontsize=8, loc="upper left")

keys = ["records", "structure", "strength", "strain", "scores", "holdout"]
for i, key in enumerate(keys):
    fig = plt.figure(figsize=(8.2, 5.1), layout="constrained")
    panel(i, fig, fig.add_gridspec(1, 1)[0])
    fig.savefig(OUT / f"{key}.svg", metadata={"Date": None})
    plt.close(fig)
fig = plt.figure(figsize=(16.4, 15.3), layout="constrained")
grid = fig.add_gridspec(3, 2, hspace=.09, wspace=.05)
for i in range(6):
    panel(i, fig, grid[i//2, i%2])
fig.savefig(OUT / "silk_learning.svg", metadata={"Date": None})
fig.savefig(OUT / "silk_learning.png", dpi=160, metadata={"Software": "L15 reproducible example"})
plt.close(fig)
print("Original workbook: SHA-256 verified")
print("108 available fracture records; 4 conditions; 4 held-out folds")
print("Model                  stress RMSE (MPa)  strain RMSE (percentage points)")
for name, label in zip(names, labels):
    stress, strain = scores[name]["condition_mean"]["RMSE"]
    print(f"{label:22s} {stress:17.3f} {strain:31.3f}")
print(f"Independent ridge solution: maximum difference {ridge_check:.2e}")
print("Six panels and complete data/results written to silk_outputs/")
